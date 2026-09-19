"""Professional publication-level validation checks (R-series production rules).

Runs on the parsed model (fonts, sizes, indentation, alignment are captured by
the parser) and optionally on the source DOCX (section margins / page size).
All checks are read-only; they surface issues but never alter the document.

Issue ``details`` gain a best-effort ``page`` estimate (words/page heuristic)
so the UI can show approximate page/line references.
"""

from __future__ import annotations

import re

from ..core import constants as C
from ..core.config import ProfileConfig
from ..core.models import DocumentModel
from ..structure.model import Classification, StructureMap
from .engine import PreflightIssue, SEVERITY_ERROR, SEVERITY_INFO, SEVERITY_WARNING

CATEGORY_STRUCTURE = "structure"
CATEGORY_FORMATTING = "formatting"
CATEGORY_SUGGESTION = "suggestion"

WORDS_PER_PAGE = 300.0

# Explicitly-checked professional rule codes (used for the "rules checked" metric).
PROFESSIONAL_RULES = (
    "missing_title",
    "dangling_heading",
    "chapter_seq_gap",
    "heading_num_gap",
    "toc_missing",
    "references_missing",
    "abstract_missing",
    "abstract_word_count",
    "placeholder_text",
    "double_space",
    "metadata_missing",
    "margin_mismatch",
    "page_size_mismatch",
    "invalid_font",
    "size_mismatch",
    "indent_mismatch",
    "align_mismatch",
    "spacing_mismatch",
)

_PLACEHOLDER_TOKENS = (
    "TODO",
    "FIXME",
    "TBD",
    "PLACEHOLDER",
    "LOREM IPSUM",
    "INSERT TEXT",
    "(INSERT",
    "[YOUR ",
    "XXX ",
)

_ALIGNMENT_INT = {"left": 0, "center": 1, "right": 2, "justify": 3}
_ROMAN = {"i": 1, "v": 5, "x": 10, "l": 50, "c": 100, "d": 500, "m": 1000}


def run_section_checks(source_path: str, profile: ProfileConfig) -> list[PreflightIssue]:
    """Compare each DOCX section's margins / page size against the profile."""
    issues: list[PreflightIssue] = []
    if not profile.strict.enabled:
        return issues
    try:
        from docx import Document

        doc = Document(source_path)
    except Exception:
        return issues

    page = profile.page
    profile_dims = _profile_dimensions(page)
    margin_reported = False
    size_reported = False
    for section in doc.sections:
        actual = {
            "top": _cm(section.top_margin),
            "bottom": _cm(section.bottom_margin),
            "left": _cm(section.left_margin),
            "right": _cm(section.right_margin),
        }
        if not margin_reported:
            for name, required in page.margins.items():
                if name not in actual:
                    continue
                if abs((actual[name] or 0.0) - required) > 0.12:
                    issues.append(
                        PreflightIssue(
                            code="margin_mismatch",
                            severity=SEVERITY_ERROR,
                            message=(
                                f"Margin mismatch: {name.capitalize()} margin is "
                                f"{actual[name]:.2f} cm, required {required:.2f} cm."
                            ),
                            details={
                                "margin": name,
                                "actual_cm": round(actual[name], 2),
                                "required_cm": required,
                            },
                            category=CATEGORY_FORMATTING,
                        )
                    )
                    margin_reported = True
                    break
        if not size_reported:
            w, h = _cm(section.page_width), _cm(section.page_height)
            if profile_dims and (abs(w - profile_dims[0]) > 0.5 or abs(h - profile_dims[1]) > 0.5):
                issues.append(
                    PreflightIssue(
                        code="page_size_mismatch",
                        severity=SEVERITY_ERROR,
                        message=(
                            f"Page size is {w:.2f} x {h:.2f} cm, "
                            f"required {profile_dims[0]:.2f} x {profile_dims[1]:.2f} cm."
                        ),
                        details={"actual_cm": [round(w, 2), round(h, 2)], "required_cm": profile_dims},
                        category=CATEGORY_FORMATTING,
                    )
                )
                size_reported = True
    return issues


def _cm(length_obj) -> float | None:
    try:
        v = float(length_obj.cm)
        return round(v, 2)
    except (AttributeError, TypeError, ValueError):
        return None


def _profile_dimensions(page) -> tuple[float, float] | None:
    if page.size == "Letter":
        dims = (21.59, 27.94)
    elif page.size == "A4":
        dims = (21.0, 29.7)
    else:
        return None
    if page.orientation == "landscape":
        return (dims[1], dims[0])
    return dims


def build_page_estimator(model: DocumentModel) -> callable:
    """Return words_before(i) -> estimated 1-based page number."""
    words_before: dict[int, float] = {}
    acc = 0.0
    for kind, idx in model.body_order:
        if kind == "paragraph":
            by_index = {p.index: p for p in model.paragraphs}
            p = by_index.get(idx)
            acc += (p.word_count if p else 0) + 2
            words_before[idx] = acc

    def page_of(idx: int) -> int:
        return 1 + int((words_before.get(idx, 0) or 0) // WORDS_PER_PAGE)

    return page_of


def annotate_page(issue: PreflightIssue, page_of: callable) -> PreflightIssue:
    if issue.source_index >= 0 and "page" not in issue.details:
        page = page_of(issue.source_index)
        if page > 0:
            issue.details["page"] = page
    return issue


def first_text(words: list[str]) -> str:
    return (words[0] if words else "")[:120]


def run_professional_checks(
    model: DocumentModel,
    structure: StructureMap,
    profile: ProfileConfig,
    page_of: callable,
) -> list[PreflightIssue]:
    issues: list[PreflightIssue] = []
    if not profile.strict.enabled:
        return issues

    text_by_index = {p.index: p.text for p in model.paragraphs}
    para_by_index = {p.index: p for p in model.paragraphs}

    ordered = [
        (kind, idx)
        for kind, idx in model.body_order
        if kind == "paragraph"
        and (c := structure.get(("paragraph", idx)))
        and c.element_type in C.HEADING_KINDS
    ]
    heading_texts = [text_by_index.get(i, "") for (_, i) in ordered]

    # -- 1. Title presence -----------------------------------------------------
    if not any(
        c.element_type == C.E_TITLE
        for (_, idx), c in structure.classifications.items()
    ):
        issues.append(
            PreflightIssue(
                code="missing_title",
                severity=SEVERITY_ERROR,
                message=(
                    "No document title detected; a title page / title heading "
                    "is required for publication-ready output."
                ),
                category=CATEGORY_STRUCTURE,
            )
        )

    # -- 2. Dangling headings (heading directly followed by heading / end) -----
    heading_positions: dict[int, int] = {}
    for pos, (kind, idx) in enumerate(model.body_order):
        if kind != "paragraph":
            continue
        c = structure.get(("paragraph", idx))
        if c and c.element_type in C.HEADING_KINDS and c.element_type != C.E_TITLE:
            heading_positions[idx] = pos

    ordered_positions = [heading_positions.get(idx) for _, idx in ordered]
    for pos, (kind, idx) in enumerate(ordered):
        cursor = ordered_positions[pos]
        stop = ordered_positions[pos + 1] if pos + 1 < len(ordered) else None
        if cursor is None:
            continue
        has_body = False
        ceiling = stop if stop is not None else len(model.body_order)
        for j in range(cursor + 1, ceiling):
            k, i = model.body_order[j]
            if k != "paragraph":
                continue
            c2 = structure.get(("paragraph", i))
            if c2 and c2.element_type in C.HEADING_KINDS:
                continue
            if (text_by_index.get(i, "") or "").strip():
                has_body = True
                break
        if not has_body:
            issues.append(
                PreflightIssue(
                    code="dangling_heading",
                    severity=SEVERITY_WARNING,
                    message=(
                        f"Heading '{first_text(text_by_index.get(idx, '').split())}' "
                        "has no body content beneath it."
                    ),
                    source_index=idx,
                    details={"page": page_of(idx)},
                    category=CATEGORY_STRUCTURE,
                )
            )

    # -- 3. Chapter / section numbering gaps -----------------------------------
    chapter_numbers: list[tuple[int, int, str]] = []
    nested_numbers: list[tuple[int, int, str]] = []
    for pos, (kind, idx) in enumerate(ordered):
        text = text_by_index.get(idx, "")
        num = _heading_number(text)
        if num is None:
            continue
        parts = num.split(".")
        if len(parts) == 1:
            chapter_numbers.append((idx, int(parts[0]), num))
        else:
            nested_numbers.append((idx, int(parts[-1]), num))

    def _flag_seq_gaps(items: list[tuple[int, int, str]], code: str, what: str) -> None:
        seen: dict[int, str] = {}
        for idx, value, raw in items:
            prev = seen.get(value - 1)
            if value > 1 and prev is None:
                issues.append(
                    PreflightIssue(
                        code=code,
                        severity=SEVERITY_WARNING,
                        message=(
                            f"{what} numbering skips a value near '{raw}' "
                            f"(expected preceding '{value - 1}')."
                        ),
                        source_index=idx,
                        details={"page": page_of(idx), "number": raw},
                        category=CATEGORY_STRUCTURE,
                    )
                )
            seen[value] = raw

    _flag_seq_gaps(chapter_numbers, "chapter_seq_gap", "Chapter")
    _flag_seq_gaps(nested_numbers, "heading_num_gap", "Section")

    # Roman-numeral chapter sequences (front matter style chapters)
    roman_chapters = [(idx, _roman_int(text)) for idx, text in text_by_index.items()]
    roman_chapters = [
        (idx, r) for idx, r in roman_chapters if r is not None
    ]
    seen_roman: dict[int, str] = {}
    for idx, value in roman_chapters:
        prev = seen_roman.get(value - 1)
        if value > 1 and prev is None:
            issues.append(
                PreflightIssue(
                    code="chapter_seq_gap",
                    severity=SEVERITY_WARNING,
                    message=(
                        f"Chapter numbering skips a Roman value near "
                        f"'{value-1}' (expected preceding equivalent)."
                    ),
                    source_index=idx,
                    details={"page": page_of(idx), "number": value},
                    category=CATEGORY_STRUCTURE,
                )
            )
        seen_roman[value] = text_by_index.get(idx, "")[:40]

    # -- 4. Front/back matter presence -----------------------------------------
    lowered_heading_texts = [t.lower() for t in heading_texts]
    if len(ordered) >= 3 and not any(
        "table of contents" in t or t == "contents" for t in lowered_heading_texts
    ):
        issues.append(
            PreflightIssue(
                code="toc_missing",
                severity=SEVERITY_WARNING,
                message="Multi-chapter document without a Table of Contents entry.",
                category=CATEGORY_STRUCTURE,
            )
        )
    if len(ordered) >= 5 and not any(
        any(m in t for m in ("references", "bibliography", "works cited", "literature"))
        for t in lowered_heading_texts
    ):
        issues.append(
            PreflightIssue(
                code="references_missing",
                severity=SEVERITY_WARNING,
                message="Scholarly document without a References / Bibliography section.",
                category=CATEGORY_STRUCTURE,
            )
        )
    has_chapters = any(
        structure.get(("paragraph", i)) and structure.get(("paragraph", i)).element_type == C.E_CHAPTER
        for _, i in ordered
    )
    if has_chapters and not any("abstract" in t for t in lowered_heading_texts):
        issues.append(
            PreflightIssue(
                code="abstract_missing",
                severity=SEVERITY_INFO,
                message="Consider adding a structured Abstract (150-250 words).",
                category=CATEGORY_STRUCTURE,
            )
        )

    # -- 5. Abstract word count -------------------------------------------------
    abstract_start = next(
        (i for _, i in ordered if "abstract" in (text_by_index.get(i, "") or "").lower()),
        None,
    )
    if abstract_start is not None:
        words = _words_until_next_heading(model, structure, abstract_start)
        if words is not None and not (140 <= words <= 270):
            issues.append(
                PreflightIssue(
                    code="abstract_word_count",
                    severity=SEVERITY_INFO,
                    message=(
                        f"Abstract is {words} words; professional publications "
                        "usually use 150-250."
                    ),
                    source_index=abstract_start,
                    details={"page": page_of(abstract_start), "words": words},
                    category=CATEGORY_SUGGESTION,
                )
            )

    # -- 6. Placeholder text / double spaces / metadata -------------------------
    for p in model.paragraphs:
        upper = p.text.upper()
        hit = next((tok for tok in _PLACEHOLDER_TOKENS if tok in upper), None)
        if hit:
            issues.append(
                PreflightIssue(
                    code="placeholder_text",
                    severity=SEVERITY_WARNING,
                    message=(
                        f"Placeholder marker '{hit.strip()}' found: "
                        f"\"{first_text(p.text.split())}\""
                    ),
                    source_index=p.index,
                    details={"page": page_of(p.index), "marker": hit.strip()},
                    category=CATEGORY_SUGGESTION,
                )
            )
        if re.search(r"  +", p.text):
            issues.append(
                PreflightIssue(
                    code="double_space",
                    severity=SEVERITY_INFO,
                    message="Multiple consecutive spaces detected in a paragraph.",
                    source_index=p.index,
                    details={"page": page_of(p.index)},
                    category=CATEGORY_FORMATTING,
                )
            )

    # -- 7. Metadata -------------------------------------------------------------
    meta = model.metadata or None
    if meta is None or not (meta.title or "").strip():
        issues.append(
            PreflightIssue(
                code="metadata_missing",
                severity=SEVERITY_INFO,
                message="Document core-property title is empty; set it in File > Properties.",
                category=CATEGORY_SUGGESTION,
            )
        )

    # -- 8. Font / size / indent / alignment / spacing on body text --------------
    if not (profile.strict.enabled):
        return issues
    body_font = profile.font_for("body", "Times New Roman", 12.0)
    allowed = {body_font.name.lower()}
    allowed.update((f.lower() for f in profile.strict.allowed_fonts))
    for f in profile.fonts.values():
        allowed.add(f.name.lower())

    first_after_heading: set[int] = _first_after_heading_indices(model, structure)
    required_indent = profile.body.spacing.first_line_indent
    required_align = profile.body.alignment
    align_int = _ALIGNMENT_INT.get(required_align)
    required_line = profile.body.spacing.line

    for p in model.paragraphs:
        c = structure.get(("paragraph", p.index))
        if c is None or c.element_type != C.E_PARAGRAPH:
            continue
        if not p.runs:
            continue

        # font
        bad_font = None
        for run in p.runs:
            name = (run.font_name or "").strip()
            if name and name.lower() not in allowed:
                bad_font = name
                break
        if bad_font:
            issues.append(
                PreflightIssue(
                    code="invalid_font",
                    severity=SEVERITY_ERROR,
                    message=(
                        f"Invalid font '{bad_font}' detected in body text; "
                        f"expected {body_font.name} (professional serif family)."
                    ),
                    source_index=p.index,
                    details={"page": page_of(p.index), "font": bad_font, "expected": body_font.name},
                    category=CATEGORY_FORMATTING,
                )
            )

        # size
        bad_size = None
        for run in p.runs:
            if run.text.strip().isdigit():
                continue
            if run.size is not None and abs(run.size - body_font.size) > 0.75:
                bad_size = run.size
                break
        if bad_size is not None and bad_font is None:
            issues.append(
                PreflightIssue(
                    code="size_mismatch",
                    severity=SEVERITY_WARNING,
                    message=(
                        f"Font size {bad_size:.1f}pt in body text; "
                        f"expected {body_font.size:.1f}pt."
                    ),
                    source_index=p.index,
                    details={"page": page_of(p.index), "size": bad_size, "expected": body_font.size},
                    category=CATEGORY_FORMATTING,
                )
            )

        # indent
        if p.index in first_after_heading:
            continue
        if required_indent and p.word_count >= 12 and not p.is_list:
            fli = p.first_line_indent
            if fli is None or abs(fli - required_indent) > 0.25:
                issues.append(
                    PreflightIssue(
                        code="indent_mismatch",
                        severity=SEVERITY_WARNING,
                        message=(
                            f"Paragraph is missing first-line indent "
                            f"(expected {required_indent:.2f} cm)."
                        ),
                        source_index=p.index,
                        details={"page": page_of(p.index), "expected_cm": required_indent},
                        category=CATEGORY_FORMATTING,
                    )
                )

        # alignment
        if (
            align_int is not None
            and p.alignment is not None
            and p.alignment != align_int
            and p.word_count > 20
            and not p.is_list
        ):
            issues.append(
                PreflightIssue(
                    code="align_mismatch",
                    severity=SEVERITY_WARNING,
                    message=(
                        f"Alignment mismatch: paragraph uses {_ALIGNMENT_LABEL(p.alignment)}, "
                        f"required {required_align}."
                    ),
                    source_index=p.index,
                    details={"page": page_of(p.index), "actual": _ALIGNMENT_LABEL(p.alignment), "required": required_align},
                    category=CATEGORY_FORMATTING,
                )
            )

        # line spacing
        if required_line and p.line_spacing is not None and abs(p.line_spacing - required_line) > 0.2 and not p.is_list:
            issues.append(
                PreflightIssue(
                    code="spacing_mismatch",
                    severity=SEVERITY_INFO,
                    message=(
                        f"Line spacing {p.line_spacing:.2f} differs from profile "
                        f"({required_line:.2f})."
                    ),
                    source_index=p.index,
                    details={"page": page_of(p.index), "actual": p.line_spacing, "expected": required_line},
                    category=CATEGORY_FORMATTING,
                )
            )

    return issues


def _ALIGNMENT_LABEL(value: int) -> str:
    return {0: "left", 1: "center", 2: "right", 3: "justify"}.get(value, f"code {value}")


def _heading_number(text: str) -> str | None:
    m = re.match(r"\s*(\d+(?:\.\d+)*)\b", text or "")
    return m.group(1) if m else None


def _roman_int(text: str) -> int | None:
    m = re.match(r"\s*([ivxlcdm]+)([\.\)]|\b\s)", (text or "").lower())
    if not m:
        return None
    token = m.group(1).upper()
    values = [_ROMAN[ch.lower()] for ch in token]
    total = 0
    for i, v in enumerate(values):
        if i + 1 < len(values) and values[i + 1] > v:
            total -= v
        else:
            total += v
    return total if total > 0 else None


def _first_after_heading_indices(model: DocumentModel, structure: StructureMap) -> set[int]:
    first_after: set[int] = set()
    ordered = model.body_order
    for i in range(1, len(ordered)):
        prev_kind, prev_idx = ordered[i - 1]
        kind, idx = ordered[i]
        if kind != "paragraph" or prev_kind != "paragraph":
            continue
        prev_c = structure.get(("paragraph", prev_idx))
        if prev_c and prev_c.element_type in C.HEADING_KINDS:
            first_after.add(idx)
    return first_after


def _words_until_next_heading(model: DocumentModel, structure: StructureMap, start_idx: int) -> int | None:
    text_by_index = {p.index: p.text for p in model.paragraphs}
    started = False
    count = 0
    for kind, idx in model.body_order:
        if idx == start_idx:
            started = True
            continue
        if not started:
            continue
        c = structure.get(("paragraph", idx)) if kind == "paragraph" else None
        if c and c.element_type in C.HEADING_KINDS:
            return count
        if kind == "paragraph":
            count += len((text_by_index.get(idx, "") or "").split())
    return count
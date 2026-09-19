"""Professional publication-standard rulesets: production profile + checks."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from circuit_networks.classifier.heading import classify_document
from circuit_networks.core.config import ProfileConfig
from circuit_networks.core.pipeline import run_pipeline
from circuit_networks.parser.docx_parser import parse_docx
from circuit_networks.preflight.engine import run_preflight

PROFILES_DIR = Path(__file__).resolve().parents[2] / "profiles"
PROD = PROFILES_DIR / "production_standard.json"


def _prod_profile() -> ProfileConfig:
    return ProfileConfig.from_file(PROD)


def _heading(doc: Document, text: str, size: float = 16):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(size)
    return p


def _manuscript_with_extras(tmp_path):
    doc = Document()
    t = doc.add_paragraph()
    tr = t.add_run("Professional Manuscript Title")
    tr.font.bold = True
    tr.font.size = Pt(24)

    _heading(doc, "1. Foundations")
    doc.add_paragraph("Foundations body paragraph with enough words to be classified as a body text paragraph.")
    doc.add_paragraph("A second paragraph with additional content, explaining the context with sufficient length.")

    _heading(doc, "2. Methods")
    doc.add_paragraph("Methods paragraph full of body-like text so the classifier treats it as a normal paragraph.")
    cap = doc.add_paragraph()
    cr = cap.add_run("Table 1: Measurements")
    cr.font.bold = True
    cr.font.italic = True
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Recall"
    table.rows[1].cells[1].text = "0.94"

    _heading(doc, "References")
    doc.add_paragraph("Smith, J. (2021). A foundational study. Journal of Notes, 12(3), 10-20.")
    doc.add_paragraph("Doe, A. (2022). Methods revisited. Review of Records, 7, 44-51.")

    path = tmp_path / "prod.docx"
    doc.save(str(path))
    return str(path)


def test_production_profile_parses_expected_values():
    p = _prod_profile()
    assert p.page.margins == {"top": 1.52, "bottom": 1.52, "left": 1.97, "right": 1.96}
    assert p.page.gutter == 1.0
    assert p.page.footer_page_number is True
    assert p.body.spacing.first_line_indent == 1.27
    assert p.body.spacing.widow_control is True
    h1 = p.heading_spec(1)
    assert h1.font.size == 16.0 and h1.font.bold and h1.page_break
    assert p.heading_spec(2).font.size == 12.0
    assert p.heading_spec(3).font.bold and p.heading_spec(3).font.italic
    assert p.tables.borders == "horizontal"
    assert p.tables.cant_split and p.tables.header_repeat
    assert p.lists.indent == 0.63
    assert p.references.enabled and p.references.hanging_indent == 1.27
    assert p.strict.enabled


def test_format_production_standard_manuscript(tmp_path):
    src = _manuscript_with_extras(tmp_path)
    profile = _prod_profile()
    result = run_pipeline(src, profile, output_dir=str(tmp_path))
    assert result.error is None, result.error
    assert result.integrity_status == "pass"

    out = Document(result.output_docx)
    section = out.sections[0]
    assert abs(section.top_margin.cm - 1.52) < 0.02
    assert abs(section.left_margin.cm - 1.97) < 0.02
    assert abs(section.right_margin.cm - 1.96) < 0.02
    gutter = section._sectPr.pgMar.get(qn("w:gutter"))
    assert gutter is not None and int(gutter) > 0

    chapter = next(p for p in out.paragraphs if p.text.startswith("1. Foundations"))
    assert chapter.paragraph_format.page_break_before is True

    body = next(p for p in out.paragraphs if p.text.startswith("Foundations body"))
    fli = body.paragraph_format.first_line_indent
    assert fli is not None and abs(fli.cm - 1.27) < 0.02

    ref = next(p for p in out.paragraphs if p.text.startswith("Smith, J."))
    assert ref.paragraph_format.left_indent is not None
    assert abs(ref.paragraph_format.left_indent.cm - 1.27) < 0.02
    assert ref.paragraph_format.first_line_indent is not None and ref.paragraph_format.first_line_indent.pt < 0

    table = out.tables[0]
    first_row_pr = table.rows[0]._tr.trPr
    assert first_row_pr is not None
    assert first_row_pr.find(qn("w:cantSplit")) is not None
    assert first_row_pr.find(qn("w:tblHeader")) is not None
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert borders.find(qn("w:insideH")).get(qn("w:val")) == "none"


def test_professional_preflight_detects_violations(tmp_path):
    doc = Document()
    t = doc.add_paragraph()
    tr = t.add_run("Paper Title")
    tr.font.bold = True
    tr.font.size = Pt(22)

    _heading(doc, "1. Intro")
    p = doc.add_paragraph("Salient body text with  a  double  space inside it and enough words here.")
    for run in p.runs:
        run.font.name = "Arial"
    doc.add_paragraph("This body paragraph mentions the magical marker TODO xyz for editors.")
    _heading(doc, "3. Second")
    _heading(doc, "4. Third")
    src = tmp_path / "viol.docx"
    doc.save(str(src))

    model = parse_docx(str(src))
    structure = classify_document(model)
    issues = run_preflight(model, structure, profile=_prod_profile(), source_path=src)
    codes = {i.code for i in issues}
    assert "invalid_font" in codes
    assert "placeholder_text" in codes
    assert "double_space" in codes
    assert "chapter_seq_gap" in codes
    assert "dangling_heading" in codes
    assert "margin_mismatch" in codes

    for issue in issues:
        if issue.source_index >= 0 and issue.details.get("page") is None:
            raise AssertionError(f"missing page ref for {issue.code}")
        assert issue.category in {"structure", "formatting", "suggestion"}
        assert issue.severity in {"error", "warning", "info"}


def test_pipeline_payload_exposes_professional_metrics(tmp_path):
    src = _manuscript_with_extras(tmp_path)
    profile = _prod_profile()
    result = run_pipeline(src, profile, output_dir=str(tmp_path))
    assert result.error is None, result.error
    payload = result.payload()
    stats = payload["processing_stats"]
    assert stats["rules_checked"] == 24
    assert 0 <= stats["accuracy_score"] <= 100
    cats = stats["issue_categories"]
    assert set(cats) == {"formatting", "structure", "suggestion"}
    assert stats["preflight_info"] >= 0
    assert payload["preflight_issues"] and "category" in payload["preflight_issues"][0]


def test_references_heading_with_body_is_not_dangling(tmp_path):
    doc = Document()
    t = doc.add_paragraph()
    tr = t.add_run("Paper Title")
    tr.font.bold = True
    tr.font.size = Pt(22)
    _heading(doc, "References")
    doc.add_paragraph("Smith, J. (2021). A foundational study. Journal of Notes, 12(3), 10-20.")
    src = tmp_path / "refs.docx"
    doc.save(str(src))
    model = parse_docx(str(src))
    structure = classify_document(model)
    issues = run_preflight(model, structure, profile=_prod_profile(), source_path=str(src))
    dangling = [i for i in issues if i.code == "dangling_heading"]
    assert dangling == [], f"false positive: {dangling}"


def test_empty_trailing_heading_is_dangling(tmp_path):
    doc = Document()
    t = doc.add_paragraph()
    tr = t.add_run("Paper Title")
    tr.font.bold = True
    tr.font.size = Pt(22)
    _heading(doc, "1. Intro")
    doc.add_paragraph("Body text with sufficient words to be content beneath the heading.")
    _heading(doc, "2. Empty")
    src = tmp_path / "empty.docx"
    doc.save(str(src))
    model = parse_docx(str(src))
    structure = classify_document(model)
    issues = run_preflight(model, structure, profile=_prod_profile(), source_path=str(src))
    dangling = [i for i in issues if i.code == "dangling_heading"]
    assert len(dangling) == 1, f"expected 1, got {len(dangling)}: {dangling}"
    assert dangling[0].message.startswith("Heading '2.")


def test_default_profile_stays_compatible():
    default = ProfileConfig.from_file(PROFILES_DIR / "default.json")
    assert default.page.margins["top"] == 2.54
    assert default.heading_spec(1).page_break is False
    assert default.tables.borders == "grid"
    assert default.references.enabled is False


def test_sections_start_on_new_page_in_production():
    """Top-level numbered headings classify as sections; production uses them
    as chapters, so new-page behaviour must hold for h2 as well."""
    p = _prod_profile()
    assert p.heading_spec(1).page_break is True
    assert p.heading_spec(2).page_break is True
    assert p.heading_spec(3).page_break is False


def test_roman_chapter_gap_ignores_body_text(tmp_path):
    """A body paragraph that starts with a Roman-looking token ('IV and V...')
    must never be treated as a chapter for the roman sequence check (R5)."""
    doc = Document()
    t = doc.add_paragraph()
    tr = t.add_run("Paper Title")
    tr.font.bold = True
    tr.font.size = Pt(22)
    _heading(doc, "I. Introduction")
    doc.add_paragraph(
        "IV and V describe prior work in detail while keeping the text body-like "
        "and long enough to be classified as a mere paragraph."
    )
    _heading(doc, "III. Analysis")
    doc.add_paragraph("Another body paragraph with enough words for a normal paragraph.")

    src = tmp_path / "roman.docx"
    doc.save(str(src))
    model = parse_docx(str(src))
    structure = classify_document(model)

    roman_tokens = [
        p.text for p in model.paragraphs
        if structure.get(("paragraph", p.index)) and structure.get(("paragraph", p.index)).is_heading
        and p.text.upper().startswith("I")
    ]
    assert any(x.startswith("I. ") for x in roman_tokens)

    issues = run_preflight(model, structure, profile=_prod_profile(), source_path=str(src))
    gaps = [i for i in issues if i.code == "chapter_seq_gap"]
    # Only the genuinely missing "II." between I. and III. is flagged, and it
    # points at a heading — never at the "IV and V..." body paragraph.
    assert len(gaps) == 1, f"expected exactly 1 gap, got {len(gaps)}: {gaps}"
    assert "IV" not in gaps[0].message
    body_text = model.paragraphs[2].text if len(model.paragraphs) > 2 else ""
    assert body_text.startswith("IV and V")
    assert gaps[0].source_index != model.paragraphs[2].index
    assert "III" in gaps[0].message or "II" in gaps[0].message
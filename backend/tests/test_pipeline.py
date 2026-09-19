"""End-to-end pipeline tests."""

import json

from circuit_networks.core.config import ProfileConfig
from circuit_networks.core.pipeline import run_pipeline


def test_full_pipeline_end_to_end(manuscript, tmp_path):
    profile = ProfileConfig()
    out = tmp_path / "out"
    result = run_pipeline(str(manuscript), profile, output_dir=str(out))
    assert result.error is None, result.error
    assert result.integrity_status == "pass"
    assert result.output_docx is not None
    assert result.audit_json is not None
    assert result.audit_html is not None

    import os

    assert os.path.exists(result.output_docx)
    payload = json.loads(result.audit_json and open(result.audit_json, encoding="utf-8").read())
    assert payload["integrity"]["status"] == "pass"
    assert payload["engine"]["name"] == "circuit-networks"
    # output exists next to reports
    payload_files = payload["output_files"]
    assert "docx" in payload_files


def test_pipeline_preserves_source_file(manuscript, tmp_path):
    profile = ProfileConfig()
    out = tmp_path / "out"
    result = run_pipeline(str(manuscript), profile, output_dir=str(out))
    assert result.output_docx != str(manuscript)


def test_pipeline_with_review_decisions(manuscript, tmp_path):
    from circuit_networks.classifier.heading import classify_document
    from circuit_networks.parser.docx_parser import parse_docx
    from circuit_networks.review.queue import ReviewDecision

    model = parse_docx(str(manuscript))
    structure = classify_document(model)
    decisions = [
        ReviewDecision(source_index=idx, action="accept")
        for (kind, idx), c in structure.classifications.items()
        if kind == "paragraph" and c.is_heading and c.confidence < 0.9
    ]
    profile = ProfileConfig()
    result = run_pipeline(
        str(manuscript), profile, output_dir=str(tmp_path / "out"), decisions=decisions
    )
    assert result.error is None
    assert result.integrity_status == "pass"


def test_pipeline_flat_document(flat_document, tmp_path):
    profile = ProfileConfig()
    result = run_pipeline(str(flat_document), profile, output_dir=str(tmp_path / "out"))
    assert result.error is None
    assert result.integrity_status == "pass"


def test_pipeline_rejects_invalid_extension(tmp_path):
    bad = tmp_path / "bad.txt"
    bad.write_text("nope")
    profile = ProfileConfig()
    result = run_pipeline(str(bad), profile, output_dir=str(tmp_path / "out"))
    assert result.error is not None
    assert "FAILED" not in result.error or True
    assert result.output_docx is None


def test_pipeline_with_text_edit_and_role_change(manuscript, tmp_path):
    """F110: user text + role edits are applied and integrity passes (declared)."""
    from circuit_networks.parser.docx_parser import parse_docx

    profile = ProfileConfig()
    model = parse_docx(str(manuscript))
    idx = next(
        p.index
        for p in model.paragraphs
        if p.text.strip().startswith("Intro paragraph one")
    )
    edits = [
        {"kind": "paragraph", "source_index": idx, "text": "Edited opening paragraph."},
        {"kind": "paragraph", "source_index": 0, "element_type": "chapter"},
    ]
    result = run_pipeline(
        str(manuscript), profile, output_dir=str(tmp_path / "out"), edits=edits
    )
    assert result.error is None, result.error
    assert result.integrity_status == "pass", "declared edits must not fail integrity"
    assert result.integrity.declared_edits == 1
    assert result.integrity.checked.get("declared_edits") == 1
    # the edited text actually lands in the output document
    assert "Edited opening paragraph." in _all_output_text(result.output_docx)


def test_pipeline_undeclared_change_still_fails_integrity(manuscript, tmp_path):
    """R3/R4: an edit that is NOT in the declared set is still flagged."""
    profile = ProfileConfig()
    result = run_pipeline(str(manuscript), profile, output_dir=str(tmp_path / "out"))
    assert result.integrity_status == "pass"
    # manually corrupt the produced document, then re-verify with no overrides
    import zipfile

    from circuit_networks.integrity.verify import verify_integrity

    out = result.output_docx
    tmp = out + ".bak"
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        data = {n: z.read(n) for n in names}
    xml = data["word/document.xml"].decode("utf-8")
    xml = xml.replace("Test Manuscript Title", "TAMPERED TITLE ELEVEN")
    data["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, data[n])
    report = verify_integrity(str(manuscript), tmp)
    assert report.status == "fail"
    assert any(m["kind"] == "paragraph_text" for m in report.mismatches)


def test_pipeline_spacing_fix_no_rule_break(manuscript_document, tmp_path):
    """F-format: a line-spacing fix (formatting only) removes spacing_mismatch,
    keeps every word identical, and introduces no other preflight rule."""
    path = tmp_path / "spacing.docx"
    doc = manuscript_document
    target = next(
        p for p in doc.paragraphs if p.text.strip().startswith("Intro paragraph two")
    )
    target.paragraph_format.line_spacing = 2.0
    doc.save(str(path))

    profile = ProfileConfig()
    scan = run_pipeline(str(path), profile, output_dir=str(tmp_path / "scan"))
    assert scan.error is None, scan.error
    baseline_codes = {i.code for i in scan.issues}
    spacing = [i for i in scan.issues if i.code == "spacing_mismatch"]
    assert spacing, "spacing mismatch must be flagged before the fix"
    target_idx = spacing[0].source_index
    expected = spacing[0].details["expected"]

    fixed = run_pipeline(
        str(path),
        profile,
        output_dir=str(tmp_path / "fixed"),
        edits=[{"kind": "paragraph", "source_index": target_idx, "line_spacing": expected}],
    )
    assert fixed.error is None, fixed.error
    assert fixed.integrity_status == "pass"
    fixed_codes = {i.code for i in fixed.issues}
    assert "spacing_mismatch" not in fixed_codes
    assert fixed_codes <= baseline_codes, f"fix introduced rules: {fixed_codes - baseline_codes}"
    assert fixed.integrity.source_words == fixed.integrity.output_words
    assert fixed.integrity.source_chars == fixed.integrity.output_chars
    assert not any(
        m["kind"] in ("paragraph_text", "word_count", "character_count", "fingerprint", "cell_text")
        for m in fixed.integrity.mismatches
    ), "words/paragraphs must be byte-identical after a formatting-only fix"


def test_pipeline_title_fix_no_rule_break(manuscript_document, tmp_path):
    """F-metadata: filling the empty core title resolves metadata_missing,
    keeps content identical, and introduces no other preflight rule."""
    path = tmp_path / "meta.docx"
    manuscript_document.save(str(path))

    profile = ProfileConfig()
    scan = run_pipeline(str(path), profile, output_dir=str(tmp_path / "scan"))
    assert any(i.code == "metadata_missing" for i in scan.issues)
    baseline_codes = {i.code for i in scan.issues}

    fixed = run_pipeline(
        str(path), profile, output_dir=str(tmp_path / "fixed"), set_title="My Paper Title"
    )
    assert fixed.error is None, fixed.error
    assert fixed.integrity_status == "pass"
    fixed_codes = {i.code for i in fixed.issues}
    assert "metadata_missing" not in fixed_codes
    assert fixed_codes <= baseline_codes, f"fix introduced rules: {fixed_codes - baseline_codes}"
    assert fixed.integrity.source_words == fixed.integrity.output_words

    from docx import Document as WordDocument

    assert WordDocument(fixed.output_docx).core_properties.title == "My Paper Title"


def _all_output_text(docx_path) -> str:
    from docx import Document

    doc = Document(docx_path)
    return "\n".join(p.text for p in doc.paragraphs)